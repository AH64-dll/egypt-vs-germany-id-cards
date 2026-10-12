"""Slides 33–37: assignment comparison table (AR + EN), key findings, references, closing."""
from common import *
from pptx.util import Inches
from lxml import etree
from pptx.oxml.ns import qn

GROUPS_EN = [("CHARACTERISTICS", 5), ("DESIGN", 3), ("SECURITY", 9)]
GROUPS_AR = [("الخصائص", 5), ("التصميم", 3), ("التأمين", 9)]

ROWS_EN = [
    ("Issuer", "Civil Status Sector, Ministry of Interior", "Federal Ministry of the Interior; made by Bundesdruckerei"),
    ("In use since", "Plastic card since 1995 (secondary source)", "Chip card since 1 Nov 2010; current design 2 May 2024"),
    ("Format", "ID-1 proportions; no published specification", "ID-1 · 8.6 × 5.4 cm"),
    ("Card body", "Plastic; polymer not officially published", "Multi-layer polycarbonate + protective lacquer"),
    ("Validity", "7 years (secondary source)", "10 years; 6 years under age 24"),
    ("Language", "Arabic only, right-to-left", "German · English · French labels"),
    ("Motifs", "Pyramids & Sphinx · pharaonic relief · Tutankhamun · eagle", "Federal eagle · Brandenburg Gate · EU flag “DE”"),
    ("Portrait", "One photo; the MLI shows birth date & emblem", "Three: inkjet photo, holographic portrait, laser image"),
    ("Security printing", "Guilloche · rainbow printing · relief “Egypt” · microprint 5×", "Multicolour guilloches · ± microprint · anti-copy · iris print"),
    ("Optically variable", "Kinegram® (Tutankhamun) · Multi Laser Image", "Identigram® · OVI logo · variable laser image"),
    ("UV features", "UV-glowing eagle · invisible UV microtext", "Eagle & lettering in rainbow UV print · UV fibres"),
    ("Tactile", "Multi Laser Image, felt with the fingers", "Tactile expiry & CAN · surface embossing"),
    ("Personalization", "Method not documented", "Inkjet photo & document no. · laser-engraved data"),
    ("Substrate & lab", "Lab marks: materials, deliberate errors, filter-only text", "Personalised security thread · security fibres"),
    ("Machine-readable", "Encrypted 2D barcode (PDF417) · factory number", "3-line MRZ · contactless chip"),
    ("Biometrics & eID", "Not documented — part of the lecture's future vision", "Face + 2 fingerprints on chip · eID with PIN"),
    ("Verification", "4 levels: public → specialists → tools → laboratory", "8 official methods incl. tilt, UV, touch, app"),
]
ROWS_AR = [
    ("الجهة المصدرة", "قطاع الأحوال المدنية · وزارة الداخلية", "وزارة الداخلية الاتحادية · تصنيع Bundesdruckerei"),
    ("منذ", "بطاقة بلاستيكية منذ 1995 (مصدر ثانوي)", "بشريحة منذ 1 نوفمبر 2010 · التصميم الحالي 2 مايو 2024"),
    ("المقاس", "بنسب ID-1 · لا مواصفة منشورة", "ID-1 · 8.6 × 5.4 سم"),
    ("جسم البطاقة", "بلاستيك · نوع البوليمر غير منشور رسميًا", "بولي كربونات متعدد الطبقات + طلاء حماية"),
    ("الصلاحية", "7 سنوات (مصدر ثانوي)", "10 سنوات · 6 سنوات قبل سن 24"),
    ("اللغة", "العربية فقط، من اليمين لليسار", "الألمانية والإنجليزية والفرنسية"),
    ("الرموز", "الأهرامات وأبو الهول · مشهد فرعوني · توت عنخ آمون · النسر", "النسر الاتحادي · بوابة براندنبورغ · علم الاتحاد الأوروبي"),
    ("الصورة الشخصية", "صورة واحدة؛ والصورة الليزرية تُظهر الميلاد والشعار", "ثلاث: حبر نفاث، هولوجرافية، صورة ليزر"),
    ("الطباعة الأمنية", "جلوشيات · طباعة تدريجية · Relief «مصر» · ميكرونية 5X", "جلوشيات متعددة الألوان · ميكرونية ± · ضد النسخ · قزحية"),
    ("المتغير بصريًا", "كينجرام (توت عنخ آمون) · الصورة الليزرية المتعددة", "Identigram® · شعار OVI · صورة ليزر متغيرة"),
    ("علامات UV", "نسر متألق · كتابة ميكرونية خفية", "نسر وكتابة بطباعة قزحية · ألياف UV"),
    ("اللمس", "الصورة الليزرية المتعددة البارزة", "تاريخ الانتهاء وCAN البارزان · بروز سطحي"),
    ("التشخيص", "الطريقة غير موثقة", "صورة ورقم الوثيقة بالحبر النفاث · حفر ليزري"),
    ("الخامة والمعمل", "علامات معملية: الخامات، أخطاء مقصودة، مرشحات", "خيط أمان مُشخَّص · ألياف أمان"),
    ("القراءة الآلية", "باركود ثنائي الأبعاد مشفر (PDF417) · رقم المصنع", "MRZ من 3 أسطر · شريحة لاتلامسية"),
    ("الحيوية والإلكترونية", "غير موثقة — ضمن رؤية المحاضرة المستقبلية", "الوجه + بصمتان على الشريحة · eID برقم PIN"),
    ("الفحص", "4 مستويات: الشخص العادي ← المتخصص ← الأدوات ← المعمل", "8 طرق رسمية منها الإمالة وUV واللمس والتطبيق"),
]

def compare_table(s, ar):
    rows_src, groups = (ROWS_AR, GROUPS_AR) if ar else (ROWS_EN, GROUPS_EN)
    hdr = ["", "", "مصر · بطاقة تحقيق الشخصية", "ألمانيا · Personalausweis"] if ar else ["", "", "EGYPT · IDENTITY VERIFICATION CARD", "GERMANY · PERSONALAUSWEIS"]
    rows = [[hdr[0], hdr[1], (hdr[2], {"color": C["org"]}), (hdr[3], {"color": C["txt"]})]]
    gi = 0; starts = []
    for g, n in groups:
        starts.append((len(rows), n, g))
        for k in range(n):
            a, e, d = rows_src[gi]; gi += 1
            rows.append(["", (a, {"font": F["ar_sb"] if ar else F["body_sb"], "color": C["txt2"]}), e, d])
    W = SW - 2 * MX; cw = [0.42, 1.75 if not ar else 1.85]; rest = (W - sum(cw)) / 2; cw += [rest, rest]
    rh = [0.36] + [0.284] * (len(rows) - 1)
    gs = table(s, MX, 1.55, W, cw, rows, rtl=ar, row_h=rh, size=10, header_size=10,
               fills=lambda r, c: C["s1"] if r == 0 else None)
    tbl = gs.table
    for r0, n, g in starts:
        c = tbl.cell(r0, 0); c.merge(tbl.cell(r0 + n - 1, 0))
        c = tbl.cell(r0, 0)
        tcPr = c._tc.get_or_add_tcPr(); tcPr.set("vert", "vert270"); tcPr.set("anchor", "ctr")
        c.fill.solid(); c.fill.fore_color.rgb = rgb({0: C["org"], 1: C["org2"], 2: C["deep"]}[len([x for x in starts if x[0] < r0])])
        fill_el = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
        for e in fill_el: tcPr.remove(e); tcPr.append(e)
        tf = c.text_frame; p = tf.paragraphs[0]
        for r in list(p.runs): r._r.getparent().remove(r._r)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = g
        from lib import _style_run
        _style_run(run, F["ar_sb"] if ar else F["disp_md"], 10 if ar else 9.5, C["ink"] if g in ("CHARACTERISTICS", "الخصائص") else C["txt"], False, 0 if ar else 150, ar)
    return gs

# ------------------------------------------------------------------ 33 AR comparison table
def s33(D):
    s = D.new(); bg(s, "bg_quiet.jpg")
    title_ar(s, "جدول المقارنة: مصر وألمانيا", w=8, size=30, y=0.72)
    text(s, MX, 0.98, 4.6, 0.3, "تكليف المحاضرة: التأمين · الخصائص · التصميم", size=10.5, font=F["ar_md"], color=C["org"], align="l")
    compare_table(s, True)
    chrome(s, "ar", 8, 33, TOTAL, "SP1 · SP2 · BMI 2024 · PAP · PAuswG · BDR · WIKI-EG · HRW")
    notes(s,
        "هذا هو جدول التكليف المطلوب في المحاضرة الثانية: مقارنة البطاقة المصرية ببطاقة هوية أخرى في علامات التأمين والخصائص والتصميم. الخصائص: الجهة المصدرة (قطاع الأحوال المدنية بوزارة الداخلية / وزارة الداخلية الاتحادية مع تصنيع Bundesdruckerei)، وبداية الاستخدام (بطاقة بلاستيكية منذ 1995 / بطاقة بشريحة منذ 2010 وتصميم حالي منذ 2 مايو 2024)، والمقاس (نسب ID-1 دون مواصفة منشورة / ID-1)، وجسم البطاقة (بلاستيك غير محدد النوع / بولي كربونات مع طلاء حماية)، والصلاحية (7 سنوات / 10 سنوات و6 قبل سن 24). التصميم: اللغة، والرموز الوطنية، وعدد مرات ظهور الصورة. التأمين: الطباعة الأمنية، والعناصر المتغيرة بصريًا، وعلامات UV، والعلامات الملموسة، وطريقة التشخيص، وأمان الخامة والعلامات المعملية، والقراءة الآلية، والقياسات الحيوية والهوية الإلكترونية، وطرق الفحص. البيانات المصرية من المحاضرة الأولى ومصادر ثانوية موسومة؛ والألمانية من دليل الوزارة والبوابة الرسمية والقانون.",
        "This is the assignment table from Lecture 2: compare the Egyptian ID card with another ID card on security features, characteristics and design. Characteristics: issuer (Civil Status Sector, Ministry of Interior / Federal Ministry of the Interior, produced by Bundesdruckerei); in use since (plastic card 1995 / chip card 2010, current design 2 May 2024); format (ID-1 proportions without a published specification / ID-1); card body (plastic, polymer unpublished / polycarbonate with protective lacquer); validity (7 years / 10 years, 6 under 24). Design: language, national motifs, how often the portrait appears. Security: security printing, optically variable elements, UV, tactile features, personalization method, substrate and laboratory marks, machine-readable elements, biometrics and eID, verification methods. Egyptian data come from Lecture 1 and labelled secondary sources; German data from the BMI guide, the official portal and the ID card act.",
        "SP2 p.27 (assignment); SP1 pp.11–21; BMI flyer May 2024; PAP; PAuswG §6; BDR; WIKI-EG (secondary: issuer name, 7 years, PDF417); HRW (1995).")
    return s

# ------------------------------------------------------------------ 34 EN comparison table
def s34(D):
    s = D.new(); bg(s, "bg_quiet.jpg")
    title_en(s, [("THE ASSIGNMENT: ", False), ("EGYPT VS GERMANY", True)], y=0.72, w=9, size=34)
    text(s, SW - MX - 4.8, 0.98, 4.8, 0.3, "SECURITY · CHARACTERISTICS · DESIGN", size=10, font=F["disp_md"], color=C["org"], spc=150, align="r")
    compare_table(s, False)
    chrome(s, "en", 8, 34, TOTAL, "SP1 · SP2 · BMI 2024 · PAP · PAuswG · BDR · WIKI-EG · HRW")
    notes(s,
        "The same comparison table in English, as required by the assignment (SP2 p.27). Read it top to bottom: characteristics first — who issues the card, since when, its format, body and validity; then design — language, national motifs and portrait; then security — the nine groups of features from printing to verification. Where the Egyptian side says 'not documented' or 'secondary source', the lecture and official sources do not state the fact; we did not fill gaps with assumptions.",
        "جدول المقارنة نفسه بالإنجليزية كما يطلب التكليف (المحاضرة الثانية ص 27). يُقرأ من أعلى لأسفل: الخصائص أولًا — الجهة المصدرة وتاريخ البدء والمقاس والجسم والصلاحية؛ ثم التصميم — اللغة والرموز والصورة؛ ثم التأمين — تسع مجموعات من الطباعة حتى طرق الفحص. حين يذكر الجانب المصري «غير موثق» أو «مصدر ثانوي» فذلك لأن المحاضرة والمصادر الرسمية لا تذكر المعلومة، ولم نملأ الفجوات بافتراضات.",
        "SP2 p.27 (assignment); SP1 pp.11–21; BMI flyer May 2024; PAP; PAuswG §6; BDR; WIKI-EG (secondary); HRW (secondary).")
    return s

# ------------------------------------------------------------------ 35 AR key findings
def s35(D):
    s = D.new(); bg(s, "bg_glow_bl.jpg")
    img(s, "cover_cards.png", -0.45, 1.6, 6.2, 6.2 / 1.5)
    X, W = 5.85, SW - MX - 5.85
    title_ar(s, "خمس نتائج", w=W, size=34)
    items = [
        ("فلسفة واحدة", "كلتاهما تطبق «سهل التحقق، صعب الاستنساخ» بعلامات موزعة على مستويات فحص متعددة."),
        ("مصر: التأمين في الطباعة", "جلوشيات وطباعة تدريجية وRelief وأحبار UV وكتابة ميكرونية، مع كينجرام وصورة ليزرية متعددة."),
        ("ألمانيا: التأمين داخل الخامة", "بولي كربونات، وبيانات بالحبر النفاث والحفر الليزري، وIdentigram، وخيط وألياف أمان."),
        ("الفجوة رقمية", "باركود مطبوع مقابل شريحة وبصمات وهوية إلكترونية؛ ورؤية المحاضرة لمصر تتجه نحو ما تطبقه ألمانيا."),
        ("الشفافية", "دليل ألماني رسمي بـ13 علامة وطرق فحصها، بينما لا تُنشر خامة البطاقة المصرية ولا طريقة تشخيصها."),
    ]
    yy = bullet_rows(s, X, 1.6, W, items, ar=True, row_h=0.7, gap=0.06, size=12.5)
    rect(s, X, yy + 0.2, W, 1.0, fill=C["deep"], alpha=0.4, line=C["org"], lw=0.75, r=0.08)
    text(s, X + 0.2, yy + 0.2, W - 0.4, 1.0, [[("الخلاصة: ", {"font": F["ar_sb"], "color": C["org"]}),
                                               ("تعتمد البطاقة المصرية على الطباعة الأمنية والفحص متعدد المستويات، وتتقدم الألمانية في دمج البيانات داخل الخامة وفي الطبقة الرقمية؛ ورؤية المحاضرة لمصر تسير في الاتجاه نفسه.", {})]],
         size=12.5, color=C["txt"], anchor="ctr", ls=1.25)
    chrome(s, "ar", 8, 35, TOTAL, "SP1 · SP2 · BMI 2024 · PAP")
    notes(s,
        "خمس نتائج: (1) فلسفة واحدة — البطاقتان تطبقان مبدأ «سهل التحقق، صعب الاستنساخ» بعلامات موزعة على مستويات فحص متعددة، من العين المجردة حتى المعمل. (2) مصر: يتركز التأمين الموثق في الطباعة الأمنية — الجلوشيات، والطباعة التدريجية، وطباعة Relief، والأحبار المتألقة، والكتابة الميكرونية الظاهرة والخفية — مع الكينجرام والصورة الليزرية المتعددة. (3) ألمانيا: يتركز التأمين داخل الخامة — جسم بولي كربونات، وبيانات بالحبر النفاث والحفر الليزري، وIdentigram، وخيط وألياف أمان. (4) الفجوة الأكبر رقمية: باركود مطبوع مشفر مقابل شريحة لاتلامسية وبصمات وهوية إلكترونية، وعناصر رؤية المحاضرة لمصر (اللاتلامسية، والقياسات الحيوية، والليزر الملموس، وUV القزحي) مطبقة أو لها ما يماثلها في ألمانيا. (5) الشفافية: لألمانيا دليل رسمي منشور بثلاث عشرة علامة وطرق فحصها، بينما المعلومات العامة عن البطاقة المصرية محدودة؛ فالخامة وطريقة التشخيص غير منشورتين. الخلاصة: تعتمد البطاقة المصرية على الطباعة الأمنية والفحص متعدد المستويات، وتتقدم الألمانية في دمج البيانات داخل الخامة وفي الطبقة الرقمية؛ ورؤية المحاضرة لمصر تسير في الاتجاه نفسه.",
        "Five findings: (1) One philosophy — both cards apply 'easy to verify, difficult to reproduce' with features spread over several verification levels, from the naked eye to the laboratory. (2) Egypt: documented security concentrates in security printing — guilloches, rainbow printing, relief printing, UV-glowing inks, visible and invisible microtext — plus the Kinegram and the Multi Laser Image. (3) Germany: security concentrates inside the material — polycarbonate body, inkjet and laser-engraved data, the Identigram, security thread and fibres. (4) The biggest gap is digital: an encrypted printed barcode versus a contactless chip with fingerprints and eID; the items of the lecture's vision for Egypt (contactless, biometrics, tactile laser, UV rainbow) are in use or have counterparts in Germany. (5) Transparency: Germany publishes an official guide to thirteen features and their checks, while public information on the Egyptian card is limited — its material and personalization method are not published. Conclusion: the Egyptian card relies on security printing and multi-level verification; the German card leads in embedding data in the material and in the digital layer; the lecture's vision for Egypt points the same way.",
        "Synthesis of slides 3–34: SP1 pp.11–21; SP2 pp.4–27; BMI flyer May 2024; PAP. Cards: Wikimedia (EG), Personalausweisportal (DE).")
    return s

# ------------------------------------------------------------------ 36 EN references
def s36(D):
    s = D.new(); bg(s, "bg_quiet.jpg")
    title_en(s, [("SOURCES & ", False), ("CREDITS", True)], y=0.72, w=8, size=36)
    sub(s, "Lectures are the primary source. Secondary sources are flagged on the slides; facts we could not verify were left out.", MX, 1.35, 11, 0.4, size=12.5)
    L = [("COURSE LECTURES", [
            ("SP1", "Security Printing, Lecture 1 — GDM5313, Dr. Rania Mohamed, ECU 2026/27 (23 pp.)"),
            ("SP2", "Security Printing, Lecture 2 (27 pp.)"),
            ("PP1", "Pre-Press Technologies, Lecture 1 — GDM5310 (27 pp.)"),
            ("PP2", "Pre-Press Technologies, Lecture 2 (35 pp.)")]),
         ("STANDARDS & LAW", [
            ("ISO", "ISO/IEC 7810 — Identification cards: physical characteristics"),
            ("PAuswG", "Personalausweisgesetz §6 — validity (gesetze-im-internet.de)")]),
         ("IMAGE CREDITS", [
            ("EG", "Wikimedia Commons, Egyptian_ID_Card.jpg, CC BY-SA 3.0 — field labels only"),
            ("EG", "Lecture photo (SP1) — personal data masked for this deck"),
            ("DE", "Specimens & macros: Personalausweisportal / BMI flyer, © BKA & Bundesdruckerei"),
            ("ART", "Diagrams, 3D renders and guilloche art: original work for this deck")])]
    R = [("GERMANY — OFFICIAL & INDUSTRY", [
            ("BMI", "BMI flyer — Security features of the German identity card, May 2024"),
            ("PAP", "Personalausweisportal (BMI) — German ID card: chip, eID, fingerprints, specimen"),
            ("BDR", "Bundesdruckerei — passport & ID card systems; ID card with online ID"),
            ("REG", "Regula Forensics — German ID card verification (secondary)"),
            ("COV", "Covestro — Makrofol® ID polycarbonate films (industry-general)"),
            ("ENG", "Plastemart / The Engineer — colour-personalised PC cards (industry-general)")]),
         ("EGYPT — SECONDARY", [
            ("WIKI-EG", "Wikipedia — Egyptian National Identity Card (validity, number structure, PDF417)"),
            ("HRW", "Human Rights Watch, 2007 (via ecoi.net) — plastic ID cards since 1995"),
            ("NAH", "Annahar FactCheck, May 2026 — no new Egyptian card design announced")])]
    for (X, groups) in [(MX, L), (SW / 2 + 0.15, R)]:
        y = 1.95; w = SW / 2 - MX - 0.15
        for g, items in groups:
            text(s, X, y, w, 0.25, g, size=10, font=F["disp_md"], color=C["org"], spc=150)
            y += 0.32
            for tag, d in items:
                text(s, X, y, 0.75, 0.3, tag, size=10, font=F["disp_md"], color=C["txt"], spc=60)
                two = len(d) > 86
                text(s, X + 0.8, y, w - 0.8, 0.5 if two else 0.3, d, size=10, color=C["txt2"], ls=1.0)
                y += 0.29 + (0.19 if two else 0)
            y += 0.16
    text(s, MX, 6.75, SW - 2 * MX, 0.25, "Full URLs, access notes and a claim-by-claim register are in the research folder (sources.md · claims-register.md · coverage-matrix.csv).",
         size=9, color=C["mute"])
    chrome(s, "en", 8, 36, TOTAL, "deck/research")
    notes(s,
        "Sources. Primary: the four course lectures (SP1, SP2, PP1, PP2). Germany: the Federal Ministry of the Interior's May 2024 flyer on security features (via bern.diplo.de), the official Personalausweisportal, the ID Card Act §6, and Bundesdruckerei; Regula is secondary; Covestro and Plastemart describe industry-general polycarbonate processing. Egypt: Wikipedia (secondary: issuer, 7-year validity, national number structure, PDF417), Human Rights Watch 2007 (plastic cards since 1995) and an Annahar fact-check (May 2026: no new design announced). Standard: ISO/IEC 7810. Images: Wikimedia Commons (Egyptian card with field labels instead of data, CC BY-SA 3.0), a lecture photo with personal data masked, German specimens and macros from the BMI and Personalausweisportal; all diagrams and renders are original. Unverified claims — for example reports of a Morpho/Safran smart-card project for Egypt — were not used on the slides.",
        "المصادر. الأساسية: المحاضرات الأربع. ألمانيا: نشرة وزارة الداخلية الاتحادية (مايو 2024) عن علامات التأمين، والبوابة الرسمية للبطاقة، وقانون بطاقة الهوية §6، وBundesdruckerei؛ وRegula مصدر ثانوي؛ وCovestro وPlastemart تصفان ممارسات صناعية عامة للبولي كربونات. مصر: ويكيبيديا (ثانوي: الجهة المصدرة، الصلاحية 7 سنوات، بنية الرقم القومي، PDF417)، وهيومن رايتس ووتش 2007 (البطاقات البلاستيكية منذ 1995)، وتدقيق النهار (مايو 2026: لا تصميم جديد). المعيار: ISO/IEC 7810. الصور: ويكيميديا (بطاقة مصرية بعناوين الحقول بدل البيانات، CC BY-SA 3.0)، وصورة من المحاضرة بعد إخفاء البيانات، ونماذج وصور مكبرة ألمانية من الوزارة والبوابة؛ وكل الرسوم أصلية. لم تُستخدم ادعاءات غير موثقة مثل تقارير مشروع بطاقة ذكية مع Morpho/Safran.",
        "research/sources.md; research/claims-register.md (E13 Morpho = unverified, notes only).")
    return s

# ------------------------------------------------------------------ 37 AR closing
def s37(D):
    s = D.new(); bg(s, "bg_cover.jpg")
    g = img(s, "guilloche_rosette2.png", -5.2, -5.6, 10.4, 10.4); pic_alpha(g, 0.6)
    img(s, "closing_cards.png", -0.35, 1.05, 7.7, 5.13)
    X, W = 7.05, SW - MX - 7.05
    eyebrow_pill(s, SW - MX - 2.25, 1.22, "نهاية العرض", ar=True, w=2.25)
    text(s, X, 1.75, W, 2.4, ["شكرًا لكم.", [("أسئلة؟", {"color": C["org"]})]], size=58, font=F["ar"], bold=True, ls=0.98)
    text(s, X, 4.15, W, 1.1, "سهلة التحقق، صعبة الاستنساخ — بطاقتان تحكيان قصة الطباعة الآمنة من جسم البطاقة حتى الشريحة.",
         size=16, font=F["ar_lt"], color=C["txt2"], ls=1.15)
    line(s, X, 5.68, SW - MX, 5.68, C["line2"], 0.75)
    cols = [("إعداد", "[اسم الطالب / الطلاب]"), ("إشراف", "د. رانيا محمد"), ("المقرر", ["الطباعة الآمنة", "GDM5313"])]
    cw = W / 3
    for i, (lab, val) in enumerate(cols):
        cx = SW - MX - (i + 1) * cw
        text(s, cx, 5.82, cw - 0.12, 0.3, lab, size=10.5, color=C["org"], font=F["ar_md"])
        text(s, cx, 6.12, cw - 0.12, 0.6, val, size=13, color=C["txt"], font=F["ar"], ls=1.0)
    text(s, MX, 6.95, 6.2, 0.3, "ظهر البطاقتين: نموذج Wikimedia (مصر) · نموذج BMI الرسمي (ألمانيا)", size=8.5, color=C["mute"], font=F["ar"], align="l")
    notes(s,
        "شكرًا لحسن الاستماع. خلاصة العرض: البطاقتان تطبقان مبدأ «سهل التحقق، صعب الاستنساخ»؛ المصرية عبر الطباعة الأمنية والفحص متعدد المستويات، والألمانية عبر دمج البيانات في جسم البولي كربونات وطبقة رقمية بشريحة وهوية إلكترونية. يسعدنا تلقي الأسئلة.",
        "Thank you for listening. In short: both cards apply 'easy to verify, difficult to reproduce' — the Egyptian card through security printing and multi-level verification, the German card by embedding data in a polycarbonate body and adding a digital layer with chip and eID. We welcome your questions.",
        "Images: Wikimedia Commons (Egyptian card back, CC BY-SA 3.0); BMI flyer May 2024 (German card back).")
    return s

SLIDES = [s33, s34, s35, s36, s37]
