"""Automated QA for the final deck: python -P build/qa.py output/Egypt-vs-Germany-ID-Cards.pptx"""
import sys, re, csv, pathlib
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from PIL import Image
import io
sys.stdout.reconfigure(encoding="utf-8")

ROOT = pathlib.Path(__file__).resolve().parent.parent
path = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "output/Egypt-vs-Germany-ID-Cards.pptx"
prs = Presentation(str(path))
if len(prs.slides) == 24:
    import runpy
    if len(sys.argv) == 1:
        sys.argv.append(str(path))
    runpy.run_path(str(pathlib.Path(__file__).parent / "qa_english.py"), run_name="__main__")
    raise SystemExit(0)
SW, SH = prs.slide_width, prs.slide_height
AR = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
LAT = re.compile(r"[A-Za-z]")
FONTS_OK = {"Barlow Condensed SemiBold", "Barlow Condensed Medium", "Barlow Condensed Light", "Barlow Condensed",
            "Barlow", "Barlow Medium", "Barlow Light", "Barlow SemiBold",
            "IBM Plex Sans Arabic", "IBM Plex Sans Arabic Medium", "IBM Plex Sans Arabic SemiBold", "IBM Plex Sans Arabic Light"}
errors, warns = [], []

def iter_text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from iter_text_frames(sh.shapes)
        if sh.has_text_frame:
            yield sh, sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield sh, cell.text_frame

slide_text, notes_text = {}, {}
for i, slide in enumerate(prs.slides, 1):
    lang = "ar" if i % 2 == 1 else "en"
    texts = []
    for sh, tf in iter_text_frames(slide.shapes):
        is_footer_src = sh.has_text_frame and abs(Emu(sh.top).inches - 7.08) < 0.02 and abs(Emu(sh.left).inches - 3.2) < 0.02
        for p in tf.paragraphs:
            t = "".join(r.text for r in p.runs)
            if not t.strip(): continue
            if not is_footer_src: texts.append(t)
            # RTL flag on Arabic paragraphs
            if AR.search(t):
                pPr = p._p.find(qn("a:pPr"))
                if pPr is None or pPr.get("rtl") != "1":
                    errors.append(f"s{i:02d} Arabic paragraph without rtl=1: {t[:40]!r}")
            for r in p.runs:
                rPr = r._r.find(qn("a:rPr"))
                lat = rPr.find(qn("a:latin")) if rPr is not None else None
                face = lat.get("typeface") if lat is not None else None
                if r.text.strip() and face not in FONTS_OK:
                    errors.append(f"s{i:02d} unexpected font {face!r} on {r.text[:30]!r}")
                if AR.search(r.text) and rPr is not None:
                    cs = rPr.find(qn("a:cs"))
                    if cs is None or "Arabic" not in cs.get("typeface", ""):
                        errors.append(f"s{i:02d} Arabic run without Arabic cs font: {r.text[:30]!r}")
    full = " ".join(texts)
    slide_text[i] = full
    nar, nlat = len(AR.findall(full)), len(LAT.findall(full))
    if lang == "en" and nar > 0:
        errors.append(f"s{i:02d} English slide contains {nar} Arabic characters")
    if lang == "ar" and nar / max(1, nar + nlat) < 0.6:
        errors.append(f"s{i:02d} Arabic slide is only {nar / max(1, nar + nlat):.0%} Arabic script")
    # notes
    nt = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    notes_text[i] = nt
    if not (AR.search(nt) and LAT.search(nt) and "Sources / المصادر" in nt):
        errors.append(f"s{i:02d} notes missing a language or the sources line")
    own = nt.split("———")[0]
    if lang == "ar" and len(AR.findall(own)) < len(LAT.findall(own)):
        errors.append(f"s{i:02d} notes: first block should be Arabic")
    if lang == "en" and AR.search(own):
        errors.append(f"s{i:02d} notes: first block should be English")
    # source tag (all but cover & closing)
    if i not in (1, 37):
        has_src = any(sh.has_text_frame and abs(Emu(sh.top).inches - 7.08) < 0.02 and abs(Emu(sh.left).inches - 3.2) < 0.02 and sh.text_frame.text.strip()
                      for sh in slide.shapes)
        if not has_src: errors.append(f"s{i:02d} missing source tag")
    # geometry & images
    for sh in slide.shapes:
        l, t, w, h = [Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height)]
        if (w <= 0 and h <= 0) or (sh.shape_type != 9 and (w <= 0 or h <= 0)):
            errors.append(f"s{i:02d} zero/negative size shape {sh.name}")
        off = l < -0.01 or t < -0.01 or l + w > Emu(SW).inches + 0.01 or t + h > Emu(SH).inches + 0.01
        if off and sh.shape_type != 13:
            warns.append(f"s{i:02d} shape extends off-slide: {sh.name} ({l:.2f},{t:.2f},{w:.2f},{h:.2f})")
        if sh.shape_type == 13:  # picture
            try:
                im = Image.open(io.BytesIO(sh.image.blob)); pw, ph = im.size
                cl, cr = sh.crop_left, sh.crop_right
                vis_px = pw * (1 - cl - cr)
                ppi = vis_px / max(w, 0.01)
                if ppi < 150:
                    warns.append(f"s{i:02d} low effective resolution {ppi:.0f} ppi: {sh.name} ({pw}x{ph}px shown {w:.2f}in wide)")
            except Exception as e:
                warns.append(f"s{i:02d} could not inspect image {sh.name}: {e}")

if len(prs.slides) != 37: errors.append(f"expected 37 slides, found {len(prs.slides)}")

# coverage matrix
cov_ok, cov_miss = 0, []
with open(ROOT / "research/coverage-matrix.csv", encoding="utf-8") as f:
    for row in csv.DictReader(f):
        kws = [k.strip().lower() for k in row["keywords"].split("|") if k.strip()]
        slides = [int(x) for x in re.findall(r"\d+", row["slides"])]
        hay = " ".join(slide_text.get(n, "") + " " + notes_text.get(n, "") for n in slides).lower()
        if any(k in hay for k in kws):
            cov_ok += 1
        else:
            where = [n for n in slide_text if any(k in (slide_text[n] + " " + notes_text[n]).lower() for k in kws)]
            cov_miss.append(f"#{row['id']} {row['concept'][:50]} — keywords {kws} not on slides {slides}; found on {where}")

print(f"QA {path.name}: {len(prs.slides)} slides")
print(f"Coverage: {cov_ok}/{cov_ok + len(cov_miss)} matrix rows matched on their mapped slides")
for m in cov_miss: print("  COVERAGE", m)
print(f"Errors: {len(errors)}")
for e in errors: print("  ERROR", e)
print(f"Warnings: {len(warns)}")
for w in warns: print("  WARN", w)
sys.exit(1 if errors or cov_miss else 0)
