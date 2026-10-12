"""Design-system toolkit on top of python-pptx: editable text (LTR + true RTL Arabic), shapes,
gradients, glows, tables, images, chrome (top nav + footer) and bilingual speaker notes."""
import re, copy, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

ROOT = pathlib.Path(__file__).resolve().parent.parent
ART, PROC = ROOT / "assets/art", ROOT / "assets/processed"

# ---------------------------------------------------------------- tokens
C = dict(ink="0A0A0A", coal="101010", s1="161616", s2="1E1E1E", s3="262626", line="2E2E2E", line2="3A3A3A",
         org="FF5A1F", org2="D63A0B", deep="8F270D", hot="FF8A4C", peach="FFB08A",
         txt="F5F3F0", txt2="A9A6A3", mute="6F6B67", black="000000")
F = dict(disp="Barlow Condensed SemiBold", disp_md="Barlow Condensed Medium", disp_lt="Barlow Condensed Light",
         disp_rg="Barlow Condensed", body="Barlow", body_md="Barlow Medium", body_lt="Barlow Light", body_sb="Barlow SemiBold",
         ar="IBM Plex Sans Arabic", ar_md="IBM Plex Sans Arabic Medium", ar_sb="IBM Plex Sans Arabic SemiBold",
         ar_lt="IBM Plex Sans Arabic Light")
SW, SH = 13.333, 7.5
MX = 0.6  # side margin
AR_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")

CHAPTERS_EN = ["CONTEXT", "MATERIALS", "LAYERS", "PRINTING", "PERSONALIZATION", "SECURITY", "FINISHING", "VERDICT"]
CHAPTERS_AR = ["السياق", "الخامات", "الطبقات", "الطباعة", "التشخيص", "التأمين", "التشطيب", "الخلاصة"]

def is_ar(s):
    return bool(AR_RE.search(s or ""))

LTR_RE = re.compile(r"[A-Za-z0-9À-ɏ][A-Za-z0-9À-ɏ .,:;/+×%°'’&§®™\-–_@#\[\]]*[A-Za-z0-9À-ɏ%°®™\]]|[A-Za-z0-9]")
LRM = "‎"

def bidi_segments(t):
    """Split text of an RTL paragraph into [(segment, is_ltr)]; LTR segments are LRM-anchored."""
    out, i = [], 0
    for m in LTR_RE.finditer(t):
        if m.start() > i: out.append((t[i:m.start()], False))
        out.append((LRM + m.group(0) + LRM, True)); i = m.end()
    if i < len(t): out.append((t[i:], False))
    return out or [(t, False)]

def rgb(h):
    return RGBColor.from_string(h)

# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def new(self):
        s = self.prs.slides.add_slide(self.blank); self.n += 1
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(C["ink"])
        return s

    def save(self, p):
        self.prs.save(p)

# ---------------------------------------------------------------- xml helpers
def _srgb(parent, hexclr, alpha=None):
    e = etree.SubElement(parent, qn("a:srgbClr"), val=hexclr)
    if alpha is not None and alpha < 1:
        etree.SubElement(e, qn("a:alpha"), val=str(int(alpha * 100000)))
    return e

def set_fill(shape, color=None, alpha=None, grad=None, angle=90):
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:gradFill", "a:noFill", "a:pattFill", "a:blipFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    if grad:
        g = etree.Element(qn("a:gradFill"), rotWithShape="1")
        gs = etree.SubElement(g, qn("a:gsLst"))
        for pos, clr, a in grad:
            st = etree.SubElement(gs, qn("a:gs"), pos=str(int(pos * 1000)))
            _srgb(st, clr, a)
        etree.SubElement(g, qn("a:lin"), ang=str(int(angle * 60000)), scaled="0")
        _insert_fill(spPr, g)
    elif color is None:
        _insert_fill(spPr, etree.Element(qn("a:noFill")))
    else:
        f = etree.Element(qn("a:solidFill")); _srgb(f, color, alpha); _insert_fill(spPr, f)

def _insert_fill(spPr, el):
    # fill must come after prstGeom/custGeom and before a:ln
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(el)
    else:
        spPr.append(el)

def set_line(shape, color=None, width=0.75, alpha=None, dash=None):
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:ln")):
        spPr.remove(e)
    ln = etree.SubElement(spPr, qn("a:ln"), w=str(int(Pt(width))))
    if color is None:
        etree.SubElement(ln, qn("a:noFill"))
    else:
        f = etree.SubElement(ln, qn("a:solidFill")); _srgb(f, color, alpha)
        if dash:
            etree.SubElement(ln, qn("a:prstDash"), val=dash)
    # effectLst must follow ln
    eff = spPr.find(qn("a:effectLst"))
    if eff is not None:
        spPr.remove(eff); spPr.append(eff)

def effects(shape, glow=None, shadow=None, soft=None):
    """glow=(radius_pt, hex, alpha); shadow=(blur_pt, dist_pt, dir_deg, hex, alpha); soft=radius_pt"""
    spPr = shape._element.spPr
    for e in spPr.findall(qn("a:effectLst")):
        spPr.remove(e)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    if glow:
        g = etree.SubElement(eff, qn("a:glow"), rad=str(int(Pt(glow[0])))); _srgb(g, glow[1], glow[2])
    if shadow:
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(Pt(shadow[0]))), dist=str(int(Pt(shadow[1]))),
                              dir=str(int(shadow[2] * 60000)), algn="ctr", rotWithShape="0")
        _srgb(sh, shadow[3], shadow[4])
    if soft:
        etree.SubElement(eff, qn("a:softEdge"), rad=str(int(Pt(soft))))

def no_shadow(shape):
    effects(shape)

# ---------------------------------------------------------------- shapes
def rect(s, x, y, w, h, fill=None, alpha=None, line=None, lw=0.75, lalpha=None, r=None, grad=None, angle=90, dash=None, name=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if r:
        shp.adjustments[0] = min(0.5, r / max(0.01, min(w, h)))
    set_fill(shp, fill, alpha, grad, angle); set_line(shp, line, lw, lalpha, dash); no_shadow(shp)
    shp.text_frame.text = ""
    if name: shp.name = name
    return shp

def oval(s, x, y, w, h, fill=None, alpha=None, line=None, lw=0.75, lalpha=None, grad=None, angle=90):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill, alpha, grad, angle); set_line(shp, line, lw, lalpha); no_shadow(shp)
    return shp

def shape(s, kind, x, y, w, h, fill=None, alpha=None, line=None, lw=0.75, rot=0, adj=None):
    shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if adj:
        for i, v in enumerate(adj): shp.adjustments[i] = v
    set_fill(shp, fill, alpha); set_line(shp, line, lw); no_shadow(shp)
    shp.rotation = rot
    return shp

def line(s, x1, y1, x2, y2, color=None, w=0.75, alpha=None, dash=None, head=None, tail=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln = c.line; ln.width = Pt(w)
    lnx = c._element.spPr.find(qn("a:ln"))
    if lnx is None:
        lnx = etree.SubElement(c._element.spPr, qn("a:ln"), w=str(int(Pt(w))))
    for e in list(lnx): lnx.remove(e)
    f = etree.SubElement(lnx, qn("a:solidFill")); _srgb(f, color or C["org"], alpha)
    if dash: etree.SubElement(lnx, qn("a:prstDash"), val=dash)
    if head: etree.SubElement(lnx, qn("a:headEnd"), type=head, w="sm", len="sm")
    if tail: etree.SubElement(lnx, qn("a:tailEnd"), type=tail, w="sm", len="sm")
    return c

def dot(s, cx, cy, d=0.09, fill=None, line_c=None):
    return oval(s, cx - d / 2, cy - d / 2, d, d, fill=fill or C["org"], line=line_c, lw=1.0)

def num_badge(s, cx, cy, n, d=0.26, fill=None, txt_c=None, size=10, font=None):
    o = oval(s, cx - d / 2, cy - d / 2, d, d, fill=fill or C["org"])
    tf = o.text_frame; _frame(tf, 0, 0, 0, 0, "ctr")
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); _style_run(r, font or F["disp"], size, txt_c or C["ink"], False, 0, is_ar(str(n)))
    return o

# ---------------------------------------------------------------- text
def _frame(tf, l=0, t=0, r=0, b=0, anchor="t", wrap=True):
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k, v in dict(lIns=l, tIns=t, rIns=r, bIns=b).items():
        bp.set(k, str(int(Inches(v))))
    bp.set("anchor", anchor); bp.set("wrap", "square" if wrap else "none")
    for e in bp.findall(qn("a:spAutoFit")) + bp.findall(qn("a:normAutofit")): bp.remove(e)

def _style_run(r, font, size, color, bold, spc, ar, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
    etree.SubElement(rPr, qn("a:latin"), typeface=font)
    etree.SubElement(rPr, qn("a:ea"), typeface=font)
    etree.SubElement(rPr, qn("a:cs"), typeface=font)
    if spc: rPr.set("spc", str(int(spc)))
    if ar:
        rPr.set("lang", "ar-EG"); rPr.set("altLang", "en-US")
    else:
        rPr.set("lang", "en-GB")

def text(s, x, y, w, h, content, size=16, color=None, font=None, bold=False, align=None, anchor="t",
         spc=0, ls=None, after=0, before=0, rtl=None, caps=False, name=None, wrap=True, italic=False):
    """content: str | list of paragraphs; paragraph: str | list of runs; run: str | (str, {overrides})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; _frame(tf, 0, 0, 0, 0, anchor, wrap)
    paras = content if isinstance(content, list) else [content]
    flat = " ".join(p if isinstance(p, str) else " ".join(r if isinstance(r, str) else r[0] for r in p) for p in paras)
    ar = is_ar(flat) if rtl is None else rtl
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = [para] if isinstance(para, str) else para
        pal = align or ("r" if ar else "l")
        p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER, "j": PP_ALIGN.JUSTIFY}[pal]
        pPr = p._p.get_or_add_pPr()
        if ar: pPr.set("rtl", "1")
        if ls:
            lnS = etree.SubElement(pPr, qn("a:lnSpc")); etree.SubElement(lnS, qn("a:spcPct"), val=str(int(ls * 100000)))
        if before:
            sb = etree.SubElement(pPr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts"), val=str(int(before * 100)))
        if after:
            sa = etree.SubElement(pPr, qn("a:spcAft")); etree.SubElement(sa, qn("a:spcPts"), val=str(int(after * 100)))
        for run in runs:
            t, o = (run, {}) if isinstance(run, str) else run
            rar = ar if o.get("rtl") is None else o["rtl"]
            fnt = o.get("font", font or (F["ar"] if rar else F["body"]))
            segs = bidi_segments(t) if ar else [(t.upper() if o.get("caps", caps) else t, not is_ar(t))]
            for seg, ltr in segs:
                r = p.add_run(); r.text = seg
                _style_run(r, fnt, o.get("size", size), o.get("color", color or C["txt"]), o.get("bold", bold),
                           o.get("spc", spc), (not ltr), o.get("italic", italic))
    if name: tb.name = name
    return tb

# ---------------------------------------------------------------- images
def img(s, path, x, y, w=None, h=None, crop=None, name=None):
    p = path if pathlib.Path(path).is_absolute() else (ART / path if (ART / path).exists() else PROC / path)
    pic = s.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if name: pic.name = name
    return pic

def img_fit(s, path, x, y, w, h, align="c"):
    """place image inside box keeping aspect (contain)."""
    from PIL import Image
    p = path if pathlib.Path(path).is_absolute() else (ART / path if (ART / path).exists() else PROC / path)
    iw, ih = Image.open(p).size; r = iw / ih
    if w / h > r:
        nw, nh = h * r, h
    else:
        nw, nh = w, w / r
    nx = x + (w - nw) / 2 if align == "c" else (x if align == "l" else x + w - nw)
    return img(s, p, nx, y + (h - nh) / 2, nw, nh)

def img_cover(s, path, x, y, w, h, fx=0.5, fy=0.5):
    """fill box (cover) cropping overflow; fx/fy focal point."""
    from PIL import Image
    p = path if pathlib.Path(path).is_absolute() else (ART / path if (ART / path).exists() else PROC / path)
    iw, ih = Image.open(p).size; r = iw / ih; br = w / h
    pic = img(s, p, x, y, w, h)
    if r > br:  # too wide -> crop left/right
        keep = br / r; extra = 1 - keep
        pic.crop_left = extra * fx; pic.crop_right = extra * (1 - fx)
    else:
        keep = r / br; extra = 1 - keep
        pic.crop_top = extra * fy; pic.crop_bottom = extra * (1 - fy)
    return pic

def round_pic(pic, radius_in=0.12):
    """turn picture geometry into rounded rectangle."""
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None: av = etree.SubElement(geom, qn("a:avLst"))
    for e in list(av): av.remove(e)
    w, h = pic.width / 914400, pic.height / 914400
    etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(50000, radius_in / min(w, h) * 100000))}")
    return pic

def circle_pic(pic):
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom")); geom.set("prst", "ellipse")
    return pic

def pic_line(pic, color, width=1.0, alpha=None):
    spPr = pic._element.spPr
    for e in spPr.findall(qn("a:ln")): spPr.remove(e)
    ln = etree.SubElement(spPr, qn("a:ln"), w=str(int(Pt(width))))
    f = etree.SubElement(ln, qn("a:solidFill")); _srgb(f, color, alpha)
    return pic

def pic_effects(pic, glow=None, shadow=None):
    spPr = pic._element.spPr
    for e in spPr.findall(qn("a:effectLst")): spPr.remove(e)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    if glow:
        g = etree.SubElement(eff, qn("a:glow"), rad=str(int(Pt(glow[0])))); _srgb(g, glow[1], glow[2])
    if shadow:
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(Pt(shadow[0]))), dist=str(int(Pt(shadow[1]))),
                              dir=str(int(shadow[2] * 60000)), algn="ctr", rotWithShape="0")
        _srgb(sh, shadow[3], shadow[4])
    return pic

def pic_alpha(pic, alpha):
    blip = pic._element.find(".//" + qn("a:blip"))
    for e in blip.findall(qn("a:alphaModFix")): blip.remove(e)
    etree.SubElement(blip, qn("a:alphaModFix"), amt=str(int(alpha * 100000)))
    return pic

# ---------------------------------------------------------------- tables
def table(s, x, y, w, col_w, rows, rtl=False, row_h=0.4, header_rows=1, size=11, header_size=10.5,
          fills=None, colors=None, fonts=None, bold_cols=(), first_col_color=None, line_c=None, valign="ctr", pad=0.08):
    """rows: list of lists of str (or (str, {overrides}))."""
    nr, nc = len(rows), len(rows[0])
    gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(row_h) if isinstance(row_h, (list, tuple)) else row_h * nr))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    if rtl: tblPr.set("rtl", "1")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None: sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, cw in enumerate(col_w): tbl.columns[i].width = Inches(cw)
    for ri in range(nr):
        tbl.rows[ri].height = Inches(row_h if not isinstance(row_h, (list, tuple)) else row_h[ri])
        for ci in range(nc):
            cell = tbl.cell(ri, ci); val = rows[ri][ci]
            t, o = (val, {}) if isinstance(val, str) else val
            hdr = ri < header_rows
            cell.margin_left = cell.margin_right = Inches(pad); cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            tcPr = cell._tc.get_or_add_tcPr(); tcPr.set("anchor", valign)
            fc = o.get("fill") or (fills(ri, ci) if callable(fills) else None)
            if fc:
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fc)
            else:
                cell.fill.background()
            # borders: bottom hairline only
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                for e in tcPr.findall(qn(tag)): tcPr.remove(e)
                ln = etree.SubElement(tcPr, qn(tag), w=str(int(Pt(0.75 if tag == "a:lnB" else 0))))
                if tag == "a:lnB":
                    f = etree.SubElement(ln, qn("a:solidFill")); _srgb(f, o.get("line", line_c or C["line"]))
                else:
                    etree.SubElement(ln, qn("a:noFill"))
            # move borders before fill per schema: lnL,lnR,lnT,lnB must precede fill
            fill_el = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
            for e in fill_el: tcPr.remove(e); tcPr.append(e)
            tf = cell.text_frame; tf.word_wrap = True
            paras = t.split("\n")
            ar = is_ar(t) or rtl
            for pi, ptxt in enumerate(paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.RIGHT if ar else PP_ALIGN.LEFT
                if o.get("align") == "c": p.alignment = PP_ALIGN.CENTER
                if ar: p._p.get_or_add_pPr().set("rtl", "1")
                segs = bidi_segments(ptxt) if ar else [(ptxt, True)]
                for seg, ltr in segs:
                    r = p.add_run(); r.text = seg
                    if hdr:
                        fnt = o.get("font", F["ar_sb"] if ar else F["disp_md"]); sz = o.get("size", header_size)
                        clr = o.get("color", (colors(ri, ci) if callable(colors) else None) or C["txt2"])
                        _style_run(r, fnt, sz, clr, False, 0 if ar else 120, not ltr)
                        if not ar: r.text = seg.upper()
                    else:
                        fnt = o.get("font", (fonts(ri, ci) if callable(fonts) else None) or (F["ar"] if ar else F["body"]))
                        clr = o.get("color", (colors(ri, ci) if callable(colors) else None) or C["txt"])
                        _style_run(r, fnt, o.get("size", size), clr, o.get("bold", ci in bold_cols), 0, not ltr)
    return gs

# ---------------------------------------------------------------- notes
def notes(s, own, other, refs):
    tf = s.notes_slide.notes_text_frame
    tf.text = own.strip() + "\n\n———\n" + other.strip() + "\n\n———\nSources / المصادر: " + refs.strip()

# ---------------------------------------------------------------- chrome
def chrome(s, lang, ch, page, total, source="", eyebrow=None, footer=True, nav=True):
    """ch: 1..8 (0 = none). Top: eyebrow + progress bars. Bottom: course, source tag, page."""
    ar = lang == "ar"
    if nav:
        label = eyebrow or ((f"{ch:02d}  —  " + (CHAPTERS_AR[ch - 1] if ar else CHAPTERS_EN[ch - 1])) if ch else "")
        if ar:
            text(s, SW - MX - 6, 0.30, 6, 0.3, [[(f"{ch:02d}" if ch else "", {"font": F["disp_md"], "color": C["org"], "rtl": False, "size": 11}),
                                                 ("   " + (CHAPTERS_AR[ch - 1] if ch else (eyebrow or "")), {"font": F["ar_md"], "color": C["txt2"], "size": 11})]],
                 align="r", rtl=True)
        else:
            text(s, MX, 0.30, 6, 0.3, [[(f"{ch:02d}" if ch else "", {"font": F["disp_md"], "color": C["org"], "spc": 150}),
                                        ("   " + (CHAPTERS_EN[ch - 1] if ch else (eyebrow or "")), {"font": F["disp_md"], "color": C["txt2"], "spc": 220})]],
                 size=11, align="l")
        # progress bars
        bw, gap = 0.20, 0.055; tot = 8 * bw + 7 * gap
        x0 = MX if ar else SW - MX - tot
        for i in range(8):
            idx = (7 - i) if ar else i
            on = (idx + 1) == ch
            rect(s, x0 + i * (bw + gap), 0.405, bw, 0.03, fill=C["org"] if on else C["line2"])
    if footer:
        y = 7.08
        course = "GDM5313 · Security Printing" if not ar else "مقرر الطباعة الآمنة · GDM5313"
        pg = f"{page:02d} / {total}"
        if ar:
            text(s, SW - MX - 4.5, y, 4.5, 0.25, course, size=9.5, color=C["mute"], font=F["ar"], align="r")
            text(s, MX, y, 1.5, 0.25, pg, size=10, color=C["txt2"], font=F["disp_md"], align="l", rtl=False, spc=100)
            if source:
                text(s, 3.2, y, 6.9, 0.25, source, size=9, color=C["mute"], font=F["body"], align="c", rtl=False)
        else:
            text(s, MX, y, 4.5, 0.25, course.upper(), size=9, color=C["mute"], font=F["disp_md"], spc=160)
            text(s, SW - MX - 1.5, y, 1.5, 0.25, pg, size=10, color=C["txt2"], font=F["disp_md"], align="r", spc=100)
            if source:
                text(s, 3.2, y, 6.9, 0.25, source, size=9, color=C["mute"], font=F["body"], align="c")

def bg(s, name):
    p = img(s, name, 0, 0, SW, SH)
    return p

def eyebrow_pill(s, x, y, label, ar=False, fill=None, color=None, w=None):
    w = w or (len(label) * (0.085 if not ar else 0.11) + 0.4)
    r = rect(s, x, y, w, 0.3, fill=fill, line=None if fill else C["org"], lw=0.75, r=0.15)
    text(s, x, y + 0.01, w, 0.28, label, size=9.5 if not ar else 10.5, color=color or C["org"],
         font=F["disp_md"] if not ar else F["ar_md"], align="c", anchor="ctr", spc=0 if ar else 160)
    return r
